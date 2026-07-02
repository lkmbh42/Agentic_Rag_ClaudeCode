"""Generate sample test documents for the ingestion pipeline.

Outputs into ./samples (or argv[1]):
  normal.pdf   text PDF                 -> text chunks
  scanned.pdf  image-only (no text)     -> OCR chunks
  table.pdf    PDF with a ruled table   -> table (+ text) chunks
  chart.pdf    PDF with an embedded PNG -> image chunk (degraded visual path)
  sample.csv   CSV                      -> table chunk
  sample.docx  Word w/ heading + table  -> text + table chunks
  sample.xlsx  spreadsheet              -> table chunk

Run inside the backend image (deps present):
  docker compose -f docker-compose.dev.yml run --rm --no-deps -v "<repo>:/srv" \
      backend sh -c "pip install --user -q reportlab && python scripts/make_samples.py"
"""

from __future__ import annotations

import io
import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "samples"
os.makedirs(OUT, exist_ok=True)


def _p(name: str) -> str:
    return os.path.join(OUT, name)


def make_normal_pdf() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(_p("normal.pdf"), pagesize=letter)
    text = c.beginText(72, 720)
    text.setFont("Helvetica", 12)
    for line in [
        "Acme Corporation Internal Policy",
        "",
        "Section 1: Remote Work",
        "Employees may work remotely up to three days per week with manager approval.",
        "Equipment is provided by the IT department upon request.",
        "",
        "Section 2: Security",
        "All staff must enable multi-factor authentication on company accounts.",
        "Confidential documents must not be shared outside the organization.",
    ]:
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()


def make_table_pdf() -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    data = [
        ["Quarter", "Revenue", "Expenses", "Profit"],
        ["Q1", "100", "60", "40"],
        ["Q2", "120", "70", "50"],
        ["Q3", "150", "90", "60"],
    ]
    t = Table(data)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    SimpleDocTemplate(_p("table.pdf"), pagesize=letter).build(
        [t]
    )


def _bar_chart_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (400, 300), "white")
    d = ImageDraw.Draw(img)
    values = [40, 50, 60, 30]
    for i, v in enumerate(values):
        x0 = 40 + i * 80
        d.rectangle([x0, 280 - v * 3, x0 + 50, 280], fill="steelblue")
    d.line([40, 280, 380, 280], fill="black", width=2)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_chart_pdf() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(_p("chart.pdf"), pagesize=letter)
    c.setFont("Helvetica", 14)
    c.drawString(72, 720, "Figure 1: Quarterly Profit (chart)")
    c.drawImage(ImageReader(io.BytesIO(_bar_chart_png())), 72, 400, width=300, height=225)
    c.showPage()
    c.save()


def make_scanned_pdf() -> None:
    """Rasterize normal.pdf into an image-only PDF (no text layer) for OCR."""
    import fitz

    src = fitz.open(_p("normal.pdf"))
    pix = src[0].get_pixmap(dpi=150)
    src.close()

    out = fitz.open()
    page = out.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    out.save(_p("scanned.pdf"))
    out.close()


def make_csv() -> None:
    with open(_p("sample.csv"), "w", encoding="utf-8") as fh:
        fh.write("name,department,role\n")
        fh.write("Alice,Engineering,Engineer\n")
        fh.write("Bob,Finance,Analyst\n")


def make_docx() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Onboarding Guide", level=1)
    d.add_paragraph("Welcome to the company. This guide covers your first week.")
    d.add_heading("Checklist", level=2)
    d.add_paragraph("Complete security training and set up your workstation.")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Task"
    table.rows[0].cells[1].text = "Owner"
    table.rows[1].cells[0].text = "Laptop setup"
    table.rows[1].cells[1].text = "IT"
    d.save(_p("sample.docx"))


def make_xlsx() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Laptop", 1500])
    ws.append(["Monitor", 300])
    wb.save(_p("sample.xlsx"))


def main() -> None:
    make_normal_pdf()
    make_table_pdf()
    make_chart_pdf()
    make_scanned_pdf()  # depends on normal.pdf
    make_csv()
    make_docx()
    make_xlsx()
    print("samples written to", os.path.abspath(OUT))
    for f in sorted(os.listdir(OUT)):
        print("  ", f, os.path.getsize(_p(f)), "bytes")


if __name__ == "__main__":
    main()
