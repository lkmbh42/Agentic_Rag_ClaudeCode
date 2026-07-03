"""Phase 2 Docling parser: typed blocks with geometry for PDF/DOCX/PPTX/XLSX,
structured tables, figure crops, OCR (deu+eng) for scanned pages.

Slow suite: real layout + TableFormer models on CPU. First run in a fresh
container downloads the Docling artifacts (dev network; prod pre-stages them).
"""

from __future__ import annotations

import io
import os

import pytest

from app.ingestion.docling_parser import parse_docling
from app.ingestion.pipeline import run_pipeline
from app.ingestion import filetype as ft
from app.models.enums import ChunkType

pytestmark = pytest.mark.docling


def _read(samples_dir: str, name: str) -> bytes:
    with open(os.path.join(samples_dir, name), "rb") as fh:
        return fh.read()


def test_pdf_text_blocks_have_pages_and_geometry(samples_dir):
    elements = parse_docling(_read(samples_dir, "normal.pdf"), ft.PDF)
    texts = [e for e in elements if e.kind == ChunkType.TEXT]
    assert texts, "no text extracted"
    assert all(e.page is not None for e in texts)
    assert any(e.bbox is not None for e in texts)


def test_pdf_table_is_structured_not_flattened(samples_dir):
    elements = parse_docling(_read(samples_dir, "table.pdf"), ft.PDF)
    tables = [e for e in elements if e.kind == ChunkType.TABLE]
    assert tables, "table not detected"
    rows = tables[0].metadata.get("rows")
    assert rows and len(rows) >= 2, "structured rows missing"
    flat = [c for row in rows for c in row]
    assert any(c for c in flat), "cells are empty"


def test_pdf_figure_carries_png_crop(samples_dir):
    elements = parse_docling(_read(samples_dir, "chart.pdf"), ft.PDF)
    figures = [e for e in elements if e.kind == ChunkType.IMAGE]
    assert figures, "embedded chart not detected as figure"
    png = figures[0].metadata.get("image_png")
    assert png and png[:8] == b"\x89PNG\r\n\x1a\n"


def test_docx_heading_becomes_section_title(samples_dir):
    elements = parse_docling(_read(samples_dir, "sample.docx"), ft.DOCX)
    kinds = {e.kind for e in elements}
    assert ChunkType.TEXT in kinds and ChunkType.TABLE in kinds
    assert any(e.section_title for e in elements if e.kind == ChunkType.TEXT)


def test_xlsx_becomes_table(samples_dir):
    elements = parse_docling(_read(samples_dir, "sample.xlsx"), ft.XLSX)
    tables = [e for e in elements if e.kind == ChunkType.TABLE]
    assert tables and tables[0].metadata.get("rows")


def test_pptx_slides_parse():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quartalsbericht"
    slide.placeholders[1].text = "Umsatz gestiegen\nKosten stabil"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    rows, cols = 3, 2
    tbl = slide2.shapes.add_table(rows, cols, Inches(1), Inches(1),
                                  Inches(6), Inches(2)).table
    for r in range(rows):
        for c in range(cols):
            tbl.cell(r, c).text = f"Z{r}S{c}"
    buf = io.BytesIO()
    prs.save(buf)

    elements = parse_docling(buf.getvalue(), ft.PPTX)
    kinds = {e.kind for e in elements}
    assert ChunkType.TEXT in kinds
    assert ChunkType.TABLE in kinds


def test_scanned_pdf_ocr_extracts_german(samples_dir):
    elements = parse_docling(_read(samples_dir, "scanned.pdf"), ft.PDF)
    text = " ".join(e.content for e in elements if e.kind == ChunkType.TEXT)
    assert len(text.strip()) > 20, "OCR produced no usable text"


# --------------------------------------------------------------- pipeline glue
def test_run_pipeline_docling_path_strips_figure_bytes(samples_dir):
    captured: list[bytes] = []

    def sink(el) -> str:
        captured.append(el.metadata["image_png"])
        return "s3://figures/test/fig.png"

    chunks, page_count = run_pipeline(_read(samples_dir, "chart.pdf"), ft.PDF,
                                      figure_sink=sink)
    assert page_count >= 1
    figure_chunks = [c for c in chunks if c.chunk_type == ChunkType.IMAGE]
    assert captured, "figure sink was not called"
    assert figure_chunks
    for c in chunks:
        assert "image_png" not in (c.source_metadata or {}), "binary leaked into metadata"
    assert figure_chunks[0].source_metadata.get("image_uri") == "s3://figures/test/fig.png"


def test_run_pipeline_table_pdf_emits_markdown_table(samples_dir):
    chunks, _ = run_pipeline(_read(samples_dir, "table.pdf"), ft.PDF)
    table_chunks = [c for c in chunks if c.chunk_type == ChunkType.TABLE]
    assert table_chunks
    assert table_chunks[0].normalized_content.startswith("Tabelle")
    assert "| ---" in table_chunks[0].normalized_content
