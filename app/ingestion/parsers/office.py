"""DOCX, XLSX, and CSV parsers."""

from __future__ import annotations

import csv
import io

from app.ingestion.types import ParsedElement
from app.models.enums import ChunkType


def _rows_to_markdown(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    return "\n".join("| " + " | ".join("" if c is None else str(c) for c in r) + " |" for r in rows)


def parse_docx(data: bytes) -> list[ParsedElement]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    elements: list[ParsedElement] = []
    section: str | None = None
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            elements.append(
                ParsedElement(kind=ChunkType.TEXT, content=body, section_title=section)
            )
        buf.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading") or style == "title":
            flush()
            section = text
        else:
            buf.append(text)
    flush()

    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        md = _rows_to_markdown(rows)
        if md:
            elements.append(ParsedElement(kind=ChunkType.TABLE, content=md, section_title=section))

    return elements or [ParsedElement(kind=ChunkType.TEXT, content="")]


def parse_xlsx(data: bytes) -> list[ParsedElement]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    elements: list[ParsedElement] = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else v) for v in row] for row in ws.iter_rows(values_only=True)]
        md = _rows_to_markdown(rows)
        if md:
            elements.append(
                ParsedElement(kind=ChunkType.TABLE, content=md, section_title=ws.title)
            )
    wb.close()
    return elements or [ParsedElement(kind=ChunkType.TABLE, content="")]


def parse_csv(data: bytes) -> list[ParsedElement]:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader]
    md = _rows_to_markdown(rows)
    return [ParsedElement(kind=ChunkType.TABLE, content=md)]
